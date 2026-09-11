from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_booking_csv(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)


def save_report_csv(
    path: Path, headers: list[str] | tuple[str, ...], records: list[dict[str, object]]
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(records)


def save_report_excel(
    path: Path, headers: list[str] | tuple[str, ...], records: list[dict[str, object]]
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet0"
    sheet.append(["COMISSÕES BOOKING × OPERA"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.append(headers)
    for record in records:
        sheet.append([record.get(header, "") for header in headers])
    sheet["A1"].font = Font(bold=True, size=12)
    sheet["A1"].alignment = Alignment(horizontal="center")
    fill = PatternFill("solid", fgColor="263238")
    for cell in sheet[2]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    for row in sheet.iter_rows(min_row=3, min_col=10, max_col=12):
        for cell in row:
            cell.number_format = '"R$" #,##0.00'
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{sheet.max_row}"
    for cells in sheet.columns:
        content_width = max(len(str(cell.value or "")) for cell in cells) + 2
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = min(
            max(content_width, 12), 48
        )
    workbook.save(path)
